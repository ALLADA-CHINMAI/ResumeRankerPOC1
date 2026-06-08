"""
xlsx_utils.py — Update candidate_data.xlsx during resume indexing.

Sheet layout:
  row 1  description text (ignored)
  row 2  column headers
  row 3+ data rows

Email is the join key between candidate_metadata (sheet 1)
and interview_history (sheet 2).
"""

import json
import logging
import os
from datetime import datetime, timezone
from typing import Optional

import openpyxl

from ResumeRankerCommon.clients import get_openai_client, OPENAI_DEPLOYMENT

logger = logging.getLogger(__name__)

CANDIDATE_DATA_PATH = os.getenv(
    "CANDIDATE_DATA_PATH",
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "candidate_data.xlsx"),
)

_EXTRACT_SYSTEM = (
    "You are a resume parser. Extract the following fields from the resume text "
    "and return valid JSON only — no markdown, no extra text:\n"
    '{\n'
    '  "full_name": "string or null",\n'
    '  "primary_email": "string or null",\n'
    '  "primary_phone": "string or null",\n'
    '  "linkedin_url": "string or null",\n'
    '  "role_family": "string or null (e.g. Software Engineer, Data Engineer, AI/ML Engineer, DevOps Engineer)",\n'
    '  "years_experience": "number or null",\n'
    '  "current_company": "string or null"\n'
    "}\n"
    "If a field is not present in the resume, use null."
)


def _extract_resume_fields(text: str) -> dict:
    """GPT-based extraction of identity and career fields from resume text."""
    try:
        resp = get_openai_client().chat.completions.create(
            model=OPENAI_DEPLOYMENT,
            messages=[
                {"role": "system", "content": _EXTRACT_SYSTEM},
                {"role": "user", "content": text[:6000]},
            ],
            temperature=0,
            max_tokens=256,
        )
        return json.loads(resp.choices[0].message.content)
    except Exception:
        logger.exception("GPT extraction failed for resume fields")
        return {}


def _interview_summary(ws_history, email: str) -> dict:
    """Aggregate interview_history rows for a given email into metadata fields."""
    headers = [c.value for c in ws_history[2]]

    rows = []
    for row in ws_history.iter_rows(min_row=3, values_only=True):
        record = dict(zip(headers, row))
        if record.get("email") == email:
            rows.append(record)

    if not rows:
        return {}

    rows.sort(key=lambda r: str(r.get("interview_date", "")), reverse=True)
    latest = rows[0]

    last_status = latest.get("result")
    last_round_type = latest.get("round_type", "")

    if last_status == "Passed" and last_round_type == "HR":
        candidate_status = "Offer Extended"
    elif last_status == "Rejected":
        candidate_status = "Rejected"
    else:
        candidate_status = "Active"

    return {
        "candidate_id":              latest.get("candidate_id"),
        "latest_application_status": last_status,
        "last_interview_round":      latest.get("round_number"),
        "total_applications":        len({r["job_role"] for r in rows if r.get("job_role")}),
        "candidate_status":          candidate_status,
    }


def update_candidate_metadata(
    doc_name: str,
    resume_text: str,
    blob_url: Optional[str] = None,
    xlsx_path: str = CANDIDATE_DATA_PATH,
) -> None:
    """
    Extract identity/career fields from a resume and upsert a row in the
    candidate_metadata sheet, enriched with interview history via email match.

    Safe to call repeatedly — re-indexing the same resume updates its row.
    Skips silently if the xlsx is missing or no email can be extracted.
    """
    if not os.path.exists(xlsx_path):
        logger.warning("candidate_data.xlsx not found at %s — skipping", xlsx_path)
        return

    fields = _extract_resume_fields(resume_text)
    email = fields.get("primary_email")
    if not email:
        logger.warning("No email extracted from %s — skipping metadata update", doc_name)
        return

    wb = openpyxl.load_workbook(xlsx_path)
    ws_meta = wb["candidate_metadata"]
    ws_hist = wb["interview_history"]

    summary = _interview_summary(ws_hist, email)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    # Headers are on row 2; data starts at row 3
    meta_headers = [c.value for c in ws_meta[2]]
    col = {h: i + 1 for i, h in enumerate(meta_headers) if h}

    # Find existing row by email or append a new one
    email_col = col.get("primary_email")
    target_row = None
    for row in ws_meta.iter_rows(min_row=3):
        if row[email_col - 1].value == email:
            target_row = row[0].row
            break
    if target_row is None:
        target_row = ws_meta.max_row + 1

    def _set(colname, value):
        idx = col.get(colname)
        if idx and value is not None:
            ws_meta.cell(row=target_row, column=idx, value=value)

    _set("candidate_id",              summary.get("candidate_id"))
    _set("full_name",                 fields.get("full_name"))
    _set("primary_email",             email)
    _set("primary_phone",             fields.get("primary_phone"))
    _set("linkedin_url",              fields.get("linkedin_url"))
    _set("role_family",               fields.get("role_family"))
    _set("years_experience",          fields.get("years_experience"))
    _set("current_company",           fields.get("current_company"))
    _set("latest_application_status", summary.get("latest_application_status"))
    _set("last_interview_round",      summary.get("last_interview_round"))
    _set("total_applications",        summary.get("total_applications"))
    _set("candidate_status",          summary.get("candidate_status"))
    _set("latest_ai_search_doc_id",   doc_name)
    _set("latest_resume_blob_url",    blob_url)

    # Preserve original created_at; always refresh updated_at
    created_at_col = col.get("created_at")
    existing_created = (
        ws_meta.cell(row=target_row, column=created_at_col).value
        if created_at_col else None
    )
    _set("created_at", existing_created or now)
    _set("updated_at", now)

    wb.save(xlsx_path)
    logger.info("candidate_metadata upserted for %s (row %d, doc=%s)", email, target_row, doc_name)
